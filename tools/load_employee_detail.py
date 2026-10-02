"""Load Factor HR's Employee Detail Report onto the Employees already on the site.

    python tools/load_employee_detail.py "Employee Detail.xlsx" --offline   # the file alone, no site
    python tools/load_employee_detail.py "Employee Detail.xlsx"             # dry run against the site
    python tools/load_employee_detail.py "Employee Detail.xlsx" --apply

    ERP_URL / ERP_KEY / ERP_SECRET in the environment, as the other tools here.
    Needs openpyxl (python -m pip install openpyxl).

**Dry run unless `--apply`**, and the dry run lists exactly what the real one
would write. The per-person list goes to `data/out/employee-detail-review.csv`,
which is gitignored with the rest of `data/`: it carries Aadhaar and account
numbers, so **nothing but counts is ever printed** — a terminal is scrollback
somebody can read over a shoulder and paste into a chat.

## What it loads

The export is 302 columns in nested sections, and the same heading appears in
several of them — `Joining Date` is the person's, then PF's, ESIC's and EPS's;
`Address` is local, then permanent. So a column is found by its section as well
as its heading (the merged cells in rows 4 and 5 say which is which), never by
the first heading that happens to match.

`FIELDS` below is the whole list. Everything not on it is left where it is:
most of the other columns are empty for all 517 people, and the rest are
either Factor HR's own switches (geofence, IMEI, reset password) or belong to
a doctype other than Employee.

## What it will not do

- **It never overwrites.** A field is written only where the site's is blank.
  A different value on the site is a conflict, listed for HR — somebody may
  have corrected it on the site since, and a bulk load is the wrong place to
  decide which is right.
- **It never writes `attendance_device_id`.** Machine codes were reconciled
  against the machines themselves (data/machine-code-*), and where they
  disagreed it was Factor HR that was stale. A difference is listed; a wrong
  device id silently drops every punch that person makes (CLAUDE.md §5).
- **It never creates an Employee.** A code the site does not hold is listed;
  `tools/build_employee_import.py` makes those, with its company mapping.
- **It never loads salary.** A salary belongs on a Salary Structure
  Assignment, against a Salary Structure nobody has built yet, and payroll is
  deliberately not started (CLAUDE.md §7). A number on Employee that no payroll
  reads is a second, disagreeing salary master.
- **It never accepts an identifier it can check and that fails the check.**
  An Aadhaar must pass its Verhoeff check digit — the export has phone numbers
  in that column; a PAN, an IFSC, a UAN and an ESIC number must have their
  shape. Factor HR's `PANNOTAVBL` is a placeholder, not a PAN.
"""

import argparse
import collections
import csv
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime
from pathlib import Path

BASE = os.environ.get("ERP_URL", "https://mannarubber.m.frappe.cloud").rstrip("/")
AUTH = "token {0}:{1}".format(os.environ.get("ERP_KEY", ""), os.environ.get("ERP_SECRET", ""))

OUT = Path("data/out/employee-detail-review.csv")

HEADER_ROW = 6
FIRST_ROW = 7

# ----------------------------------------------------------- the checks ---

# Verhoeff: the check digit every Aadhaar carries. One mistyped or transposed
# digit fails it, and so does a phone number pasted into the column.
_D = [[0, 1, 2, 3, 4, 5, 6, 7, 8, 9], [1, 2, 3, 4, 0, 6, 7, 8, 9, 5], [2, 3, 4, 0, 1, 7, 8, 9, 5, 6],
      [3, 4, 0, 1, 2, 8, 9, 5, 6, 7], [4, 0, 1, 2, 3, 9, 5, 6, 7, 8], [5, 9, 8, 7, 6, 0, 4, 3, 2, 1],
      [6, 5, 9, 8, 7, 1, 0, 4, 3, 2], [7, 6, 5, 9, 8, 2, 1, 0, 4, 3], [8, 7, 6, 5, 9, 3, 2, 1, 0, 4],
      [9, 8, 7, 6, 5, 4, 3, 2, 1, 0]]
_P = [[0, 1, 2, 3, 4, 5, 6, 7, 8, 9], [1, 5, 7, 6, 2, 8, 3, 0, 9, 4], [5, 8, 0, 3, 7, 9, 6, 1, 4, 2],
      [8, 9, 1, 6, 0, 4, 3, 5, 2, 7], [9, 4, 5, 3, 1, 2, 6, 8, 7, 0], [4, 2, 8, 6, 5, 7, 3, 9, 0, 1],
      [2, 7, 9, 3, 8, 0, 6, 4, 1, 5], [7, 0, 4, 6, 9, 1, 3, 2, 5, 8]]


# What a check returns for a cell that says "nothing here" in Factor HR's own
# words — a 0, an `Unknown`, a 1900 date. Skipped, not refused: a refusal is a
# value that claims to be something and fails the check.
EMPTY = object()


def verhoeff_ok(digits):
	c = 0
	for i, ch in enumerate(reversed(digits)):
		c = _D[c][_P[i % 8][int(ch)]]
	return c == 0


def digits_only(v):
	return re.sub(r"[\s+-]", "", str(v or ""))


def aadhaar(v):
	"""Twelve digits, never starting 0 or 1, and the check digit holds."""
	d = digits_only(v)
	return d if re.fullmatch(r"[2-9]\d{11}", d) and verhoeff_ok(d) else None


def pan(v):
	p = str(v or "").strip().upper()
	return p if re.fullmatch(r"[A-Z]{5}\d{4}[A-Z]", p) else None


def ifsc(v):
	s = str(v or "").strip().upper()
	return s if re.fullmatch(r"[A-Z]{4}0[A-Z0-9]{6}", s) else None


def uan(v):
	d = digits_only(v)
	return d if re.fullmatch(r"\d{12}", d) else None


def esic(v):
	d = digits_only(v)
	return d if re.fullmatch(r"\d{10}", d) else None


def mobile(v):
	"""An Indian mobile number. Factor HR's `1111111111` is a placeholder."""
	d = digits_only(v)
	if d.startswith("91") and len(d) == 12:
		d = d[2:]
	return d if re.fullmatch(r"[6-9]\d{9}", d) and len(set(d)) > 1 else None


def account_no(v):
	d = re.sub(r"\s", "", str(v or ""))
	return d if re.fullmatch(r"[0-9A-Za-z-]{6,34}", d) else None


def email(v):
	s = str(v or "").strip()
	return s if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", s) else None


def text(v):
	s = " ".join(str(v or "").split())
	return s or None


def iso(v):
	"""A cell's date as `YYYY-MM-DD`. Factor HR writes 02-Jan-1900 for "none"."""
	if isinstance(v, datetime):
		v = v.date()
	if isinstance(v, date):
		return v.isoformat() if v.year > 1900 else EMPTY
	m = re.fullmatch(r"(\d{2})-([A-Za-z]{3})-(\d{4})", str(v or "").strip())
	if m:
		try:
			d = datetime.strptime(m.group(0), "%d-%b-%Y").date()
		except ValueError:
			return None
		return d.isoformat() if d.year > 1900 else EMPTY
	return None


def positive_int(v):
	try:
		n = int(float(v))
	except (TypeError, ValueError):
		return None
	return n if n > 0 else EMPTY


SALUTATION = {"mr": "Mr", "mrs": "Mrs", "miss": "Miss", "ms": "Ms", "dr": "Dr"}
MARITAL = {"married": "Married", "unmarried": "Single", "single": "Single", "divorced": "Divorced",
           "widowed": "Widowed"}


def salutation(v):
	return SALUTATION.get(str(v or "").strip().rstrip(".").lower())


def marital(v):
	"""`Unknown` is Factor HR's blank, and a blank is not a status."""
	s = str(v or "").strip().lower()
	return EMPTY if s == "unknown" else MARITAL.get(s)


def reason(v):
	"""`BetterOpportunity` → `Better Opportunity`."""
	s = text(v)
	return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", s) if s else None


def salary_mode(v):
	return "Bank" if str(v or "").strip().lower() == "banktransfer" else None


# ------------------------------------------------------------ the fields ---

# (Employee field, section, heading, check). Section "" is the person's own
# columns ahead of the first labelled block. Headings are Factor HR's.
FIELDS = [
	("salutation", "", "Title", salutation),
	("date_of_birth", "", "Birth Date", iso),
	("custom_group_joining_date", "", "Group Date of Joining", iso),
	("custom_confirmation_date", "", "Confirmation Date", iso),
	("relieving_date", "", "Leaving Date", iso),
	("custom_last_working_date", "", "Employee Last working date", iso),
	("cell_number", "", "Mobile No", mobile),
	("company_email", "", "Email", email),
	("personal_email", "", "Personal Email", email),
	("custom_pan_no", "", "Pan No", pan),
	("custom_nationality", "", "Nationality", text),
	("marital_status", "", "Marital Status", marital),
	("custom_father_name", "", "Father Name", text),
	("custom_mother_name", "", "Mother Name", text),
	("custom_spouse_name", "", "Spouse Name", text),
	("custom_religion", "", "Religion", text),
	("blood_group", "", "Blood Group", text),
	("custom_probation_period_days", "", "Probation Period In Days", positive_int),
	("bank_name", "Bank Detail", "Bank Name", text),
	("bank_ac_no", "Bank Detail", "Account No", account_no),
	("custom_ifsc", "Bank Detail", "IFSC Code", ifsc),
	("salary_mode", "Bank Detail", "Payment Method", salary_mode),
	("provident_fund_account", "PF Detail", "PF No", text),
	("custom_uan", "PF Detail", "UA No", uan),
	("custom_esic_number", "ESIC Detail", "ESIC No", esic),
	("custom_aadhaar_number", "Aadhaar Card Number Details", "Aadhaar No", aadhaar),
	("resignation_letter_date", "Separation Detail", "Separation Submission Date", iso),
	("reason_for_leaving", "Separation Detail", "Reason for Leaving", reason),
]

# Read and compared, never written — see "What it will not do".
COMPARE = [
	("attendance_device_id", "", "Machine Code", text),
	("date_of_joining", "", "Joining Date", iso),
]

# The two addresses are assembled from seven columns each.
ADDRESS_PARTS = ["Address", "Address 2", "Address 3", "City", "State", "Pincode", "Country"]


# ------------------------------------------------------------ the sheet ---

def column_map(headers, spans):
	"""{(section, heading): column index} for the header row.

	`spans` is [(first col, last col, label)] from the merged cells above the
	header, 0-based and inclusive. A column inside a span belongs to that
	span's label; one inside none belongs to section "", the person's own. The
	first column found for a (section, heading) wins — so `Joining Date` is the
	person's and not PF's, and the Emergency Contact block's second contact,
	which repeats the first one's headings, is never read.

	**The two address merges are shorter than the addresses.** Factor HR's
	Permanent Address merge covers four of its seven columns, leaving State,
	City and Pincode under no label at all; so an address span runs on over
	every following column whose heading is an address part."""
	names = [" ".join(str(h or "").split()) for h in headers]
	grown = []
	for a, b, label in spans:
		if label.endswith("Address"):
			while b + 1 < len(names) and (names[b + 1] in ADDRESS_PARTS or names[b + 1] == "City Type"):
				b += 1
		grown.append((a, b, label))
	out = {}
	for i, h in enumerate(names):
		if not h:
			continue
		inside = [label for a, b, label in grown if a <= i <= b]
		for label in inside or [""]:
			out.setdefault((label, h), i)
	return out


def address(row, cols, section):
	parts = []
	for p in ADDRESS_PARTS:
		i = cols.get((section, p))
		v = text(row[i]) if i is not None and i < len(row) else None
		if v:
			parts.append(v)
	return ", ".join(parts) or None


def person(row, cols):
	"""One person's row as {field: checked value}, and [(field, why)] refused."""
	values, refused = {}, []
	for field, section, heading, check in FIELDS + COMPARE:
		i = cols.get((section, heading))
		raw = row[i] if i is not None and i < len(row) else None
		if raw in (None, "") or str(raw).strip() in ("", "PANNOTAVBL"):
			continue
		v = check(raw)
		if v is None:
			refused.append((field, "failed its check"))
		elif v is not EMPTY:
			values[field] = v
	local = address(row, cols, "Local Address")
	if local:
		values["current_address"] = local
	i = cols.get(("", "Permanent Same As Local Address"))
	same_as_local = i is not None and i < len(row) and str(row[i] or "").strip().upper() == "YES"
	permanent = local if same_as_local else address(row, cols, "Permanent Address")
	if permanent:
		values["permanent_address"] = permanent
	return values, refused


def read_sheet(path):
	"""[{code, name, values, refused, manager, salary}], and the continuation rows skipped."""
	try:
		from openpyxl import load_workbook
	except ImportError:
		sys.exit("Needs openpyxl: python -m pip install openpyxl")
	ws = load_workbook(str(path), data_only=True).worksheets[0]
	spans = [(r.min_col - 1, r.max_col - 1, " ".join(str(ws.cell(r.min_row, r.min_col).value or "").split()))
	         for r in ws.merged_cells.ranges if r.min_row == HEADER_ROW - 1]
	headers = [c.value for c in ws[HEADER_ROW]]
	cols = column_map(headers, spans)
	for need in [("", "S#"), ("", "Emp Code"), ("", "Full Name")]:
		if need not in cols:
			sys.exit("No '%s' column in the header row of %s - is this the Employee Detail Report?" % (need[1], path))
	manager = cols.get(("", "Reporting Manager"))
	salary = [i for (sec, _), i in cols.items() if sec == "Salary Master Detail"]
	people, continuation = [], 0
	for row in ws.iter_rows(min_row=FIRST_ROW, values_only=True):
		if not any(v not in (None, "") for v in row):
			continue
		# A row with no serial number is the person above's second transfer,
		# promotion or bank line — Factor HR prints one line per child row.
		if row[cols[("", "S#")]] in (None, ""):
			continuation += 1
			continue
		values, refused = person(row, cols)
		people.append({
			"code": text(row[cols[("", "Emp Code")]]) or "",
			"name": text(row[cols[("", "Full Name")]]) or "",
			"values": values,
			"refused": refused,
			"manager": text(row[manager]) if manager is not None else None,
			"salary": any(row[i] not in (None, "") for i in salary),
		})
	return people, continuation


# --------------------------------------------------------- the matching ---

def code_candidates(code):
	"""'HPT-099(BMS)' may be held as that, as 'HPT-099 (BMS)', or as plain 'HPT-099'."""
	out = [code]
	m = re.fullmatch(r"(\S+?)\s*(\([A-Z]+\))", code)
	if m:
		out += [m.group(1) + m.group(2), m.group(1) + " " + m.group(2), m.group(1)]
	return list(dict.fromkeys(out))


def _words(name):
	return {w for w in re.split(r"[^a-z]+", (name or "").lower()) if len(w) > 1}


def names_agree(a, b):
	"""Share one real word. Loose on purpose: 'SOUJATH P.I' and 'Soujath P I' are one person."""
	return bool(_words(a) & _words(b))


def match(people, employees):
	"""{code: employee} for what can be placed, and {code: why not} for the rest.

	A code nobody holds, or more than one holds, or held by somebody whose name
	shares no word with the report's, is refused — never resolved by taking the
	first. One person's Aadhaar on another's record is not a mistake anybody
	downstream would catch."""
	by_number = collections.defaultdict(list)
	for e in employees:
		if e.get("employee_number"):
			by_number[e["employee_number"].strip().upper()].append(e)
	found, refused = {}, {}
	for p in people:
		hits = []
		for c in code_candidates(p["code"]):
			hits = by_number.get(c.upper(), [])
			if hits:
				break
		if not hits:
			refused[p["code"]] = "not on the site"
		elif len(hits) > 1:
			refused[p["code"]] = "%d employees hold this code" % len(hits)
		elif not names_agree(p["name"], hits[0].get("employee_name")):
			refused[p["code"]] = "the site's name for this code shares no word with the report's"
		else:
			found[p["code"]] = hits[0]
	return found, refused


def manager_of(raw, by_code):
	"""'HPT-072 - AJITH S' → that employee's `name`, or None."""
	if not raw:
		return None
	code = raw.split(" - ")[0].strip()
	for c in code_candidates(code):
		e = by_code.get(c.upper())
		if e:
			return e["name"]
	return None


def blank(v):
	return v is None or (isinstance(v, str) and not v.strip())


def same(a, b):
	return str(a).strip().lower() == str(b).strip().lower()


def plan(values, site, writable):
	"""({field: value} to write, [(field, file value, site value)] conflicts, [field] already right).

	Writes only into a blank. `writable` is the set of fields the site's
	Employee actually has — a field the app defines but the site has not been
	given is skipped rather than sent, because Frappe answers an unknown field
	on a PUT by quietly dropping it, and a load that reports success for data
	it threw away is worse than one that says it could not."""
	fills, conflicts, same_already = {}, [], []
	for field, v in values.items():
		if field not in writable:
			continue
		have = site.get(field)
		if blank(have):
			fills[field] = v
		elif same(have, v):
			same_already.append(field)
		else:
			conflicts.append((field, v, have))
	return fills, conflicts, same_already


# The statuses under which somebody is already out of attendance and payroll.
GONE = ("Left", "Inactive")

# Fields that only mean anything on somebody who is no longer working.
LEAVING = ("relieving_date", "custom_last_working_date", "resignation_letter_date", "reason_for_leaving")


def hold_back(values, employee, by_name, salutations):
	"""(values to go ahead, [(field, why)] held back) — what the site should not be given.

	- **A leaving date on somebody the site still counts as working.** Factor HR
	  and the site disagree about whether this person works here, and a
	  relieving date is what takes somebody out of attendance and payroll. That
	  is HR's call, one person at a time, not a bulk load's. Inactive and Left
	  are both out already, so their dates go in; changing Inactive to Left is
	  a separate decision this does not take.
	- **Somebody still working reporting to a manager who has left.** ERPNext
	  then refuses to save the manager at all ("still reporting to"), so one
	  stale line in Factor HR would block every later edit of that manager.
	- **A salutation the site has no record of.** It is a Link, and a bad link
	  fails the whole save, taking the person's Aadhaar and bank with it."""
	keep, held = dict(values), []
	status = employee.get("status") or "no status"
	if status not in GONE:
		for f in LEAVING:
			if f in keep:
				del keep[f]
				held.append((f, "left in Factor HR, %s on the site" % status))
	boss = keep.get("reports_to")
	if boss and status != "Left" and (by_name.get(boss) or {}).get("status") == "Left":
		del keep["reports_to"]
		held.append(("reports_to", "the manager has left"))
	if "salutation" in keep and keep["salutation"] not in salutations:
		del keep["salutation"]
		held.append(("salutation", "not a Salutation on the site"))
	return keep, held


# ------------------------------------------------------------------ site ---

def call(method, path, body=None):
	data = json.dumps(body).encode("utf-8") if body is not None else None
	r = urllib.request.Request(BASE + path, data=data, method=method, headers={
		"Authorization": AUTH, "Accept": "application/json", "Content-Type": "application/json"})
	with urllib.request.urlopen(r, timeout=60) as f:
		return json.loads(f.read().decode("utf-8"))


def listing(doctype, fields):
	q = {"doctype": doctype, "fields": json.dumps(fields), "limit_page_length": 0}
	return call("GET", "/api/method/frappe.client.get_list?" + urllib.parse.urlencode(q)).get("message") or []


def refused_by_site(exc):
	body = exc.read().decode("utf-8", "replace")
	try:
		msgs = json.loads(json.loads(body).get("_server_messages") or "[]")
		return " ".join(json.loads(m).get("message", "") for m in msgs) or body[:300]
	except Exception:
		return body[:300]


def site_fields(wanted):
	"""Which of `wanted` the site's Employee has.

	**Not read off a record.** The REST API serialises a document without its
	null fields, so one Employee's keys are the fields *that* person happens to
	have filled — the first run of this tool took that for the schema and called
	cell_number and bank_name missing. The doctype's own meta is asked instead;
	a key without desk access is answered one field at a time, by a list read
	that names the field, which the site refuses for a field it has not got."""
	try:
		meta = call("GET", "/api/method/frappe.desk.form.load.getdoctype?" +
		            urllib.parse.urlencode({"doctype": "Employee"}))
		have = {f.get("fieldname") for d in meta.get("docs") or [] if d.get("name") == "Employee"
		        for f in d.get("fields") or []}
		if have:
			return {f for f in wanted if f in have}
	except urllib.error.HTTPError:
		pass
	out = set()
	for f in wanted:
		try:
			q = {"doctype": "Employee", "fields": json.dumps(["name", f]), "limit_page_length": 1}
			call("GET", "/api/method/frappe.client.get_list?" + urllib.parse.urlencode(q))
			out.add(f)
		except urllib.error.HTTPError:
			pass
	return out


# ------------------------------------------------------------------- run ---

def write_review(rows):
	OUT.parent.mkdir(parents=True, exist_ok=True)
	with open(OUT, "w", newline="", encoding="utf-8") as f:
		w = csv.writer(f)
		w.writerow(["code", "name", "employee", "field", "action", "file value", "site value"])
		w.writerows(rows)


def main():
	ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
	ap.add_argument("source", type=Path, help="Factor HR's Employee Detail Report, .xlsx")
	ap.add_argument("--offline", action="store_true", help="read and check the file only; no site")
	ap.add_argument("--apply", action="store_true", help="write the fills (dry run without it)")
	args = ap.parse_args()

	people, continuation = read_sheet(args.source)
	print("%d people in %s, and %d continuation rows (second transfer, promotion or bank lines) not loaded."
	      % (len(people), args.source.name, continuation))
	good = collections.Counter(f for p in people for f in p["values"])
	bad = collections.Counter(f for p in people for f, _ in p["refused"])
	print("\n%-30s %6s %8s" % ("field", "values", "refused"))
	for field in [f for f, *_ in FIELDS + COMPARE] + ["current_address", "permanent_address"]:
		if good[field] or bad[field]:
			print("%-30s %6d %8d" % (field, good[field], bad[field]))
	salaried = sum(1 for p in people if p["salary"])
	print("\nSalary figures for %d people are in the file and are not loaded - see this tool's docstring." % salaried)

	review = [[p["code"], p["name"], "", f, "refused: " + why, "", ""] for p in people for f, why in p["refused"]]

	if args.offline:
		write_review(review)
		print("\nOffline: nothing was read from or sent to the site. Per-person refusals: %s" % OUT)
		return

	if not os.environ.get("ERP_KEY"):
		sys.exit("Set ERP_URL, ERP_KEY and ERP_SECRET in the environment first (or run with --offline).")
	# A key is ASCII. Anything else dies inside http.client as a latin-1
	# UnicodeEncodeError that names neither the key nor the fix.
	if not (os.environ.get("ERP_KEY", "") + os.environ.get("ERP_SECRET", "")).isascii():
		sys.exit("ERP_KEY / ERP_SECRET hold something that is not a key - the '...' placeholder, "
		         "or a character pasted from a document. Put the real pair from ERPNext -> your User "
		         "-> API Access -> Generate Keys in their place.")

	targets = [f for f, *_ in FIELDS] + ["current_address", "permanent_address", "reports_to"]
	writable = site_fields(targets + [f for f, *_ in COMPARE])
	missing = [f for f in targets if f not in writable]
	if missing:
		print("\nNot on the site's Employee, so not loaded: %s" % ", ".join(missing))
		print("tools/create_custom_fields.py --apply puts the custom_ ones there.")
	compare = [f for f, *_ in COMPARE if f in writable]
	employees = listing("Employee", ["name", "employee_number", "employee_name", "status"] +
	                    [f for f in targets + compare if f in writable])
	by_code = {e["employee_number"].strip().upper(): e for e in employees if e.get("employee_number")}
	by_name = {e["name"]: e for e in employees}
	try:
		salutations = {r["name"] for r in listing("Salutation", ["name"])}
	except urllib.error.HTTPError:
		salutations = set()
	found, not_placed = match(people, employees)

	totals = collections.Counter()
	sent = failed = 0
	for p in people:
		e = found.get(p["code"])
		if not e:
			review.append([p["code"], p["name"], "", "", "skipped: " + not_placed[p["code"]], "", ""])
			continue
		values = dict(p["values"])
		boss = manager_of(p["manager"], by_code)
		if boss and boss != e["name"]:
			values["reports_to"] = boss
		for f in compare:
			if f in values and not blank(e.get(f)) and not same(e.get(f), values[f]):
				review.append([p["code"], p["name"], e["name"], f, "differs (never written)", values[f], e.get(f)])
		values, held = hold_back({k: v for k, v in values.items() if k not in compare}, e, by_name, salutations)
		for f, why in held:
			totals["held"] += 1
			review.append([p["code"], p["name"], e["name"], f, "held back: " + why, "", e.get(f)])
		fills, conflicts, _ = plan(values, e, writable)
		for f, v, have in conflicts:
			totals["conflict"] += 1
			review.append([p["code"], p["name"], e["name"], f, "conflict (site kept)", v, have])
		for f, v in fills.items():
			totals[f] += 1
			review.append([p["code"], p["name"], e["name"], f, "fill" if args.apply else "would fill", v, ""])
		if fills and args.apply:
			try:
				call("PUT", "/api/resource/Employee/" + urllib.parse.quote(e["name"]), fills)
				sent += 1
			except urllib.error.HTTPError as exc:
				failed += 1
				review.append([p["code"], p["name"], e["name"], "", "site refused: " + refused_by_site(exc), "", ""])

	write_review(review)
	print("\n%d matched to an employee on the site; %d not placed (%s)." % (
		len(found), len(not_placed),
		", ".join("%d %s" % (n, why) for why, n in collections.Counter(not_placed.values()).most_common())))
	print("\n%-30s %6s" % ("field", "fills"))
	for f in targets:
		if totals[f]:
			print("%-30s %6d" % (f, totals[f]))
	print("\n%d conflicts: the site already holds a different value, and keeps it." % totals["conflict"])
	print("%d held back: leaving dates on people the site calls Active, people reporting to a "
	      "manager who has left, salutations the site has not got." % totals["held"])
	if args.apply:
		print("%d employees updated, %d refused by the site." % (sent, failed))
	else:
		print("Dry run: nothing was written. Run again with --apply.")
	print("Per person, field by field: %s" % OUT)


if __name__ == "__main__":
	main()
